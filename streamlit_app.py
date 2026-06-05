import streamlit as st
import language_tool_python
import time
import fitz
from docx import Document
import io

st.set_page_config(page_title="Autocorrect — NLP", page_icon="✨", layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@400;700&family=Lato:wght@300;400;700&family=Dancing+Script:wght@600&display=swap');

* { margin: 0; padding: 0; box-sizing: border-box; }

.stApp {
    background: linear-gradient(160deg, #fdf6f0 0%, #fef9f5 40%, #f8f0ea 100%);
    color: #3d2b1f;
}

section[data-testid="stSidebar"] { display: none; }

div[data-testid="stTextArea"] textarea {
    background: rgba(255,255,255,0.8) !important;
    border: 2px solid #e8d5c4 !important;
    border-radius: 20px !important;
    color: #3d2b1f !important;
    font-family: 'Lato', sans-serif !important;
    font-size: 15px !important;
    padding: 16px !important;
}

div[data-testid="stButton"] button {
    background: linear-gradient(135deg, #c4956a, #a67c5b) !important;
    color: white !important;
    border: none !important;
    border-radius: 50px !important;
    padding: 14px 40px !important;
    font-family: 'Lato', sans-serif !important;
    font-size: 13px !important;
    font-weight: 700 !important;
    letter-spacing: 3px !important;
    width: 100% !important;
}

div[data-testid="stFileUploader"] {
    background: white !important;
    border: 2px dashed #e8d5c4 !important;
    border-radius: 20px !important;
    padding: 20px !important;
}

.block-container { padding: 2rem 3rem !important; max-width: 1200px !important; }

div[data-testid="stTabs"] button {
    font-family: 'Lato', sans-serif !important;
    font-weight: 700 !important;
    letter-spacing: 2px !important;
    font-size: 12px !important;
    color: #8a6a5a !important;
}
</style>
""", unsafe_allow_html=True)

# HERO
st.markdown('<div style="text-align:center;padding:50px 20px 20px;"><div style="font-family:Lato,sans-serif;font-size:10px;letter-spacing:4px;color:#c4956a;font-weight:700;">✦ NLP + DOCUMENT PROCESSING PROJECT ✦</div></div>', unsafe_allow_html=True)
st.markdown('<div style="text-align:center;"><div style="font-family:Dancing Script,cursive;font-size:22px;color:#c4956a;">your words, perfected</div><h1 style="font-family:Playfair Display,serif;font-size:60px;font-weight:700;color:#3d2b1f;">Auto<span style="color:#c4956a;font-style:italic;">correct</span></h1><p style="font-family:Lato,sans-serif;font-size:15px;color:#8a6a5a;max-width:600px;margin:0 auto 20px;">An intelligent NLP-powered spell correction tool with Document Processing — supports live text, PDF and Word documents</p></div>', unsafe_allow_html=True)
st.markdown('<div style="display:flex;justify-content:center;gap:10px;flex-wrap:wrap;margin-bottom:30px;"><span style="background:#fff5ef;border:1px solid #e8d5c4;border-radius:50px;padding:6px 18px;font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:2px;">NLP</span><span style="background:#fff5ef;border:1px solid #e8d5c4;border-radius:50px;padding:6px 18px;font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:2px;">DOCUMENT PROCESSING</span><span style="background:#fff5ef;border:1px solid #e8d5c4;border-radius:50px;padding:6px 18px;font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:2px;">PDF</span><span style="background:#fff5ef;border:1px solid #e8d5c4;border-radius:50px;padding:6px 18px;font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:2px;">WORD DOCS</span><span style="background:#fff5ef;border:1px solid #e8d5c4;border-radius:50px;padding:6px 18px;font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:2px;">PYTHON</span></div>', unsafe_allow_html=True)
st.markdown('<div style="width:60px;height:2px;background:linear-gradient(90deg,transparent,#c4956a,transparent);margin:0 auto 40px;"></div>', unsafe_allow_html=True)

# HOW IT WORKS
st.markdown("""
<div style="display:grid;grid-template-columns:repeat(4,1fr);gap:16px;margin-bottom:50px;">
    <div style="background:white;border-radius:24px;padding:24px;box-shadow:0 4px 20px rgba(196,149,106,0.1);border:1px solid #f0e0d0;text-align:center;">
        <div style="font-size:28px;margin-bottom:10px;">🧠</div>
        <div style="font-family:'Playfair Display',serif;font-size:16px;color:#3d2b1f;font-weight:700;margin-bottom:6px;">Smart NLP</div>
        <div style="font-family:'Lato',sans-serif;font-size:12px;color:#8a6a5a;line-height:1.6;">LanguageTool detects spelling and grammar errors intelligently</div>
    </div>
    <div style="background:white;border-radius:24px;padding:24px;box-shadow:0 4px 20px rgba(196,149,106,0.1);border:1px solid #f0e0d0;text-align:center;">
        <div style="font-size:28px;margin-bottom:10px;">📚</div>
        <div style="font-family:'Playfair Display',serif;font-size:16px;color:#3d2b1f;font-weight:700;margin-bottom:6px;">Auto Proper Nouns</div>
        <div style="font-family:'Lato',sans-serif;font-size:12px;color:#8a6a5a;line-height:1.6;">Automatically skips names and technical words</div>
    </div>
    <div style="background:white;border-radius:24px;padding:24px;box-shadow:0 4px 20px rgba(196,149,106,0.1);border:1px solid #f0e0d0;text-align:center;">
        <div style="font-size:28px;margin-bottom:10px;">📄</div>
        <div style="font-family:'Playfair Display',serif;font-size:16px;color:#3d2b1f;font-weight:700;margin-bottom:6px;">PDF Processing</div>
        <div style="font-family:'Lato',sans-serif;font-size:12px;color:#8a6a5a;line-height:1.6;">Extract text from PDFs and correct spelling</div>
    </div>
    <div style="background:white;border-radius:24px;padding:24px;box-shadow:0 4px 20px rgba(196,149,106,0.1);border:1px solid #f0e0d0;text-align:center;">
        <div style="font-size:28px;margin-bottom:10px;">📝</div>
        <div style="font-family:'Playfair Display',serif;font-size:16px;color:#3d2b1f;font-weight:700;margin-bottom:6px;">Word Docs</div>
        <div style="font-family:'Lato',sans-serif;font-size:12px;color:#8a6a5a;line-height:1.6;">Process .docx files and download corrected version</div>
    </div>
</div>
""", unsafe_allow_html=True)

# BACKEND
@st.cache_resource
def load_tool():
    return language_tool_python.LanguageTool('en-US')

import spacy
nlp = spacy.load("en_core_web_sm")

def autocorrect_text(text, tool):
    # Detect proper nouns using spaCy NER
    doc = nlp(text)
    
    # Collect all named entities to skip
    skip_words = set()
    for ent in doc.ents:
        for word in ent.text.split():
            skip_words.add(word.lower())
    
    # Also skip words tagged as proper nouns
    for token in doc:
        if token.pos_ == 'PROPN':
            skip_words.add(token.text.lower())

    # Get corrections from LanguageTool
    matches = tool.check(text)
    
    # Filter out matches that involve proper nouns
    filtered_matches = []
    for match in matches:
        wrong = match.matched_text.strip()
        if wrong.lower() not in skip_words:
            filtered_matches.append(match)
    
    # Build corrections dict
    corrections = {}
    for match in filtered_matches:
        if match.replacements:
            wrong = match.matched_text
            right = match.replacements[0]
            if wrong and right and wrong != right:
                corrections[wrong] = right

    corrected = language_tool_python.utils.correct(text, filtered_matches)
    return corrected, corrections

def extract_pdf_text(file):
    doc = fitz.open(stream=file.read(), filetype="pdf")
    text = ""
    for page in doc:
        text += page.get_text()
    return text

def extract_docx_text(file):
    doc = Document(file)
    return '\n'.join([para.text for para in doc.paragraphs])

def create_corrected_docx(corrected_text):
    doc = Document()
    doc.add_heading('Corrected Document', 0)
    for para in corrected_text.split('\n'):
        if para.strip():
            doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

with st.spinner("Loading NLP engine... please wait ⏳"):
    tool = load_tool()

# TABS
st.markdown('<div style="text-align:center;margin-bottom:24px;"><span style="font-family:Lato,sans-serif;font-size:10px;letter-spacing:4px;color:#c4956a;font-weight:700;">✦ TRY IT ✦</span><h2 style="font-family:Playfair Display,serif;font-size:36px;color:#3d2b1f;margin-top:8px;">Choose your input method</h2></div>', unsafe_allow_html=True)

tab1, tab2, tab3 = st.tabs(["✏️  LIVE TEXT", "📄  PDF DOCUMENT", "📝  WORD DOCUMENT"])

# TAB 1: LIVE TEXT
with tab1:
    col_left, col_right = st.columns(2, gap="large")

    with col_left:
        st.markdown('<p style="font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:3px;margin-bottom:8px;">✏️ YOUR TEXT</p>', unsafe_allow_html=True)
        user_text = st.text_area(label="", placeholder="Typ yur tekst heer and wach the magic hapn...", height=220, key="input")
        word_count = len(user_text.split()) if user_text.strip() else 0

        st.markdown(f"""
        <div style="display:flex;gap:12px;margin:12px 0;">
            <div style="flex:1;background:white;border-radius:16px;padding:14px;text-align:center;border:1px solid #f0e0d0;">
                <div style="font-family:'Playfair Display',serif;font-size:28px;color:#c4956a;">{word_count}</div>
                <div style="font-family:'Lato',sans-serif;font-size:10px;color:#8a6a5a;letter-spacing:2px;font-weight:700;">WORDS</div>
            </div>
            <div style="flex:1;background:white;border-radius:16px;padding:14px;text-align:center;border:1px solid #f0e0d0;">
                <div style="font-family:'Playfair Display',serif;font-size:28px;color:#c4956a;">{len(user_text)}</div>
                <div style="font-family:'Lato',sans-serif;font-size:10px;color:#8a6a5a;letter-spacing:2px;font-weight:700;">CHARS</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        run = st.button("✨ CORRECT MY TEXT", key="run_text")

    with col_right:
        st.markdown('<p style="font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:3px;margin-bottom:8px;">✅ CORRECTED OUTPUT</p>', unsafe_allow_html=True)

        if run and user_text.strip():
            with st.spinner("Checking spelling and grammar..."):
                corrected, corrections = autocorrect_text(user_text, tool)

            st.markdown(f'<div style="background:white;border-radius:20px;padding:24px;border:2px solid #e8d5c4;font-family:Lato,sans-serif;font-size:15px;line-height:1.9;color:#3d2b1f;min-height:220px;max-height:400px;overflow-y:auto;">{corrected}</div>', unsafe_allow_html=True)

           

            if corrections:
                st.markdown('<p style="font-family:Lato,sans-serif;font-size:10px;color:#e07070;font-weight:700;letter-spacing:3px;margin:8px 0;">🔴 CORRECTIONS</p>', unsafe_allow_html=True)
                for orig, corr in corrections.items():
                    st.markdown(f'<div style="background:white;border-radius:12px;padding:12px 18px;margin:6px 0;border:1px solid #f0e0d0;display:flex;justify-content:space-between;align-items:center;"><span style="font-family:Lato,sans-serif;color:#e07070;font-weight:700;text-decoration:line-through;">{orig}</span><span style="color:#c4956a;font-size:18px;">→</span><span style="font-family:Lato,sans-serif;color:#70a870;font-weight:700;">{corr}</span></div>', unsafe_allow_html=True)
            else:
                st.success("✅ Perfect spelling!")
        else:
            st.markdown('<div style="background:white;border-radius:20px;padding:40px;text-align:center;border:2px dashed #e8d5c4;min-height:220px;display:flex;flex-direction:column;justify-content:center;align-items:center;"><div style="font-size:40px;margin-bottom:12px;">✨</div><div style="font-family:Playfair Display,serif;font-size:20px;color:#c4b5a5;">Your corrected text will appear here</div></div>', unsafe_allow_html=True)

# TAB 2: PDF
with tab2:
    st.markdown('<p style="font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:3px;margin-bottom:16px;">📄 UPLOAD YOUR PDF</p>', unsafe_allow_html=True)

    pdf_file = st.file_uploader("", type=["pdf"], key="pdf_upload")

    if pdf_file:
        with st.spinner("Extracting text from PDF..."):
            pdf_text = extract_pdf_text(pdf_file)

        st.markdown(f'<div style="background:white;border-radius:16px;padding:20px;border:1px solid #f0e0d0;margin-bottom:16px;max-height:300px;overflow-y:auto;"><p style="font-family:Lato,sans-serif;font-size:10px;color:#c4956a;font-weight:700;letter-spacing:2px;margin-bottom:10px;">📖 EXTRACTED TEXT ({len(pdf_text.split())} words)</p><p style="font-family:Lato,sans-serif;font-size:13px;color:#3d2b1f;line-height:1.8;white-space:pre-wrap;">{pdf_text}</p></div>', unsafe_allow_html=True)

        if st.button("✨ CORRECT PDF TEXT", key="run_pdf"):
            with st.spinner("Running NLP spell correction..."):
                corrected, corrections = autocorrect_text(pdf_text, tool)

            st.markdown(f'<p style="font-family:Lato,sans-serif;font-size:10px;color:#70b870;font-weight:700;letter-spacing:2px;margin-bottom:10px;">✅ {len(corrections)} CORRECTIONS MADE</p>', unsafe_allow_html=True)
            st.markdown(f'<div style="background:white;border-radius:16px;padding:20px;border:1px solid #e8f5e8;margin-bottom:16px;max-height:300px;overflow-y:auto;"><p style="font-family:Lato,sans-serif;font-size:13px;color:#3d2b1f;line-height:1.8;white-space:pre-wrap;">{corrected}</p></div>', unsafe_allow_html=True)

            if corrections:
                st.markdown('<p style="font-family:Lato,sans-serif;font-size:10px;color:#e07070;font-weight:700;letter-spacing:3px;margin:8px 0;">🔴 CORRECTIONS MADE</p>', unsafe_allow_html=True)
                for orig, corr in corrections.items():
                    st.markdown(f'<div style="background:white;border-radius:12px;padding:10px 16px;margin:4px 0;border:1px solid #f0e0d0;display:flex;justify-content:space-between;align-items:center;"><span style="font-family:Lato,sans-serif;color:#e07070;font-weight:700;text-decoration:line-through;">{orig}</span><span style="color:#c4956a;font-size:16px;">→</span><span style="font-family:Lato,sans-serif;color:#70a870;font-weight:700;">{corr}</span></div>', unsafe_allow_html=True)

            corrected_doc = create_corrected_docx(corrected)
            st.download_button(
                label="⬇️ DOWNLOAD CORRECTED DOCUMENT",
                data=corrected_doc,
                file_name="corrected_document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

# TAB 3: WORD DOC
with tab3:
    st.markdown('<p style="font-family:Lato,sans-serif;font-size:11px;color:#c4956a;font-weight:700;letter-spacing:3px;margin-bottom:16px;">📝 UPLOAD YOUR WORD DOCUMENT</p>', unsafe_allow_html=True)

    docx_file = st.file_uploader("", type=["docx"], key="docx_upload")

    if docx_file:
        with st.spinner("Extracting text from Word document..."):
            docx_text = extract_docx_text(docx_file)

        st.markdown(f'<div style="background:white;border-radius:16px;padding:20px;border:1px solid #f0e0d0;margin-bottom:16px;max-height:300px;overflow-y:auto;"><p style="font-family:Lato,sans-serif;font-size:10px;color:#c4956a;font-weight:700;letter-spacing:2px;margin-bottom:10px;">📖 EXTRACTED TEXT ({len(docx_text.split())} words)</p><p style="font-family:Lato,sans-serif;font-size:13px;color:#3d2b1f;line-height:1.8;white-space:pre-wrap;">{docx_text}</p></div>', unsafe_allow_html=True)

        if st.button("✨ CORRECT DOCUMENT", key="run_docx"):
            with st.spinner("Running NLP spell correction..."):
                corrected, corrections = autocorrect_text(docx_text, tool)

            st.markdown(f'<p style="font-family:Lato,sans-serif;font-size:10px;color:#70b870;font-weight:700;letter-spacing:2px;margin-bottom:10px;">✅ {len(corrections)} CORRECTIONS MADE</p>', unsafe_allow_html=True)
            st.markdown(f'<div style="background:white;border-radius:16px;padding:20px;border:1px solid #e8f5e8;margin-bottom:16px;max-height:300px;overflow-y:auto;"><p style="font-family:Lato,sans-serif;font-size:13px;color:#3d2b1f;line-height:1.8;white-space:pre-wrap;">{corrected}</p></div>', unsafe_allow_html=True)

            if corrections:
                st.markdown('<p style="font-family:Lato,sans-serif;font-size:10px;color:#e07070;font-weight:700;letter-spacing:3px;margin:8px 0;">🔴 CORRECTIONS MADE</p>', unsafe_allow_html=True)
                for orig, corr in corrections.items():
                    st.markdown(f'<div style="background:white;border-radius:12px;padding:10px 16px;margin:4px 0;border:1px solid #f0e0d0;display:flex;justify-content:space-between;align-items:center;"><span style="font-family:Lato,sans-serif;color:#e07070;font-weight:700;text-decoration:line-through;">{orig}</span><span style="color:#c4956a;font-size:16px;">→</span><span style="font-family:Lato,sans-serif;color:#70a870;font-weight:700;">{corr}</span></div>', unsafe_allow_html=True)

            corrected_doc = create_corrected_docx(corrected)
            st.download_button(
                label="⬇️ DOWNLOAD CORRECTED DOCUMENT",
                data=corrected_doc,
                file_name="corrected_document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

# FOOTER
st.markdown("<br><br>", unsafe_allow_html=True)
st.markdown('<div style="text-align:center;padding:20px;border-top:1px solid #f0e0d0;"><div style="font-family:Dancing Script,cursive;font-size:20px;color:#c4b5a5;">made with love & language ✦</div><div style="font-family:Lato,sans-serif;font-size:10px;color:#c4b5a5;letter-spacing:3px;margin-top:8px;">NLP + DOCUMENT PROCESSING · SUMITA SAHA · KIIT UNIVERSITY · <a href="https://github.com/Sumita312/Autocorrect_Tool" style="color:#c4956a;">GITHUB</a></div></div>', unsafe_allow_html=True)